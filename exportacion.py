import pandas as pd

from configuracion import ARCHIVO_SALIDA, ZONAS
from utilidades import documentos_dataframe, nombre_hoja


def exportar_excel(datos_por_coleccion, auditoria, resumen_zonas, colecciones_relacionadas):
    print("\n")
    print("=" * 80)
    print("PREPARANDO EXCEL")
    print("=" * 80)

    for coleccion, documentos in datos_por_coleccion.items():
        unicos = {}
        for documento in documentos:
            if "_id" in documento:
                clave = (coleccion, str(documento["_id"]), documento.get("_Zona"))
            else:
                clave = (coleccion, str(documento))
            unicos[str(clave)] = documento
        datos_por_coleccion[coleccion] = list(unicos.values())

    dfs = {}
    for coleccion, documentos in datos_por_coleccion.items():
        if not documentos:
            continue
        try:
            df = documentos_dataframe(documentos)
            if not df.empty:
                dfs[nombre_hoja(coleccion)] = df
        except Exception as error:
            print(f"Error procesando {coleccion}: {error}")

    df_resumen_zonas = pd.DataFrame(resumen_zonas)
    df_auditoria = pd.DataFrame(auditoria)
    df_zonas = pd.DataFrame([
        {"Zona": nombre, "Zona_ID": zona_id}
        for nombre, zona_id in ZONAS.items()
    ])

    catalogo_colecciones = [
        {
            "Coleccion": relacion["coleccion"],
            "Campos_Zona": ", ".join(relacion["campos_zona"]),
            "Campos_Sucursal": ", ".join(relacion["campos_sucursal"]),
            "Campos_Sucursales": ", ".join(relacion["campos_sucursales"]),
            "Campos_Fecha": ", ".join(relacion.get("campos_fecha", [])),
        }
        for relacion in colecciones_relacionadas
    ]
    df_catalogo = pd.DataFrame(catalogo_colecciones)

    print(f"\nGenerando:\n{ARCHIVO_SALIDA}")
    with pd.ExcelWriter(ARCHIVO_SALIDA, engine="openpyxl") as writer:
        df_resumen_zonas.to_excel(writer, sheet_name="Resumen_Zonas", index=False)
        df_zonas.to_excel(writer, sheet_name="Zonas", index=False)
        df_auditoria.to_excel(writer, sheet_name="Auditoria", index=False)
        df_catalogo.to_excel(writer, sheet_name="Colecciones", index=False)

        max_filas = 1_000_000
        for nombre, df in dfs.items():
            nombre_sheet = nombre_hoja(nombre)
            if len(df) <= max_filas:
                df.to_excel(writer, sheet_name=nombre_sheet, index=False)
                continue

            partes = (len(df) // max_filas) + 1
            for indice in range(partes):
                inicio = indice * max_filas
                fin = min((indice + 1) * max_filas, len(df))
                parte = df.iloc[inicio:fin]
                nombre_parte = f"{nombre_sheet[:25]}_{indice + 1}"
                parte.to_excel(
                    writer,
                    sheet_name=nombre_hoja(nombre_parte),
                    index=False,
                )

    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        workbook = load_workbook(ARCHIVO_SALIDA)
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            for cell in worksheet[1]:
                cell.font = Font(bold=True)

            if worksheet.max_row > 1:
                worksheet.auto_filter.ref = worksheet.dimensions

            for column_cells in worksheet.columns:
                max_length = 0
                column_letter = get_column_letter(column_cells[0].column)
                for cell in column_cells[:100]:
                    try:
                        max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                worksheet.column_dimensions[column_letter].width = min(
                    max(max_length + 2, 10), 40
                )

        workbook.save(ARCHIVO_SALIDA)
    except Exception as error:
        print(f"Advertencia: no se pudo aplicar formato Excel: {error}")

    return dfs